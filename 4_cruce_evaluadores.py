import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.table import Table, TableStyleInfo

# Configuración de rutas
carpeta_raiz = "C:/report_download/manejo_reportes/"
carpeta_descargas = "C:/report_download/descargas/"
archivos_asignaciones = {
    "CCM": f"{carpeta_raiz}/CCM/ASIGNACIONES.xlsx",
    "PRR": f"{carpeta_raiz}/PRR/ASIGNACIONES.xlsx",
    "CCM-ESP": f"{carpeta_raiz}/CCM/ASIGNACIONES.xlsx",  # Usa la misma base de CCM
    "SOL": f"{carpeta_raiz}/SOL/ASIGNACIONES.xlsx"
}
consolidados = {
    "CCM": f"{carpeta_descargas}/CCM/Consolidado_CCM.xlsx",
    "PRR": f"{carpeta_descargas}/PRR/Consolidado_PRR.xlsx",
    "CCM-ESP": f"{carpeta_descargas}/CCM-ESP/Consolidado_CCM-ESP.xlsx",
    "SOL": f"{carpeta_descargas}/SOL/Consolidado_SOL.xlsx"
}
hojas_evaluadores = {
    "CCM": "CONSOLIDADO_CCM_X_EVAL",
    "PRR": "CONSOLIDADO_PRR_X_EVAL",
    "CCM-ESP": "CONSOLIDADO_CCM_X_EVAL",  # Usa la hoja de evaluadores de CCM
    "SOL": "CONSOLIDADO_SOL_X_EVAL"
}

# Función para procesar y cruzar los datos
def procesar_cruce():
    for tipo in consolidados.keys():
        try:
            consolidado_path = consolidados[tipo]
            asignaciones_path = archivos_asignaciones[tipo]
            hoja_evaluador = hojas_evaluadores[tipo]

            # Generar el nombre del archivo cruzado
            output_path = consolidado_path.replace(".xlsx", "_CRUZADO.xlsx")

            # Verificar si el archivo cruzado ya existe
            if os.path.exists(output_path):
                print(f"El archivo {output_path} ya existe. Saltando...")
                continue

            if not os.path.exists(consolidado_path) or not os.path.exists(asignaciones_path):
                print(f"Archivos faltantes para {tipo}. Saltando...")
                continue

            print(f"Procesando: {consolidado_path} con {asignaciones_path} ({hoja_evaluador})")

            # Leer consolidado y asignaciones
            df_consolidado = pd.read_excel(consolidado_path)
            df_asignaciones = pd.read_excel(asignaciones_path, sheet_name=hoja_evaluador)

            # Normalizar columnas
            df_consolidado.columns = df_consolidado.columns.str.strip()
            df_asignaciones.columns = df_asignaciones.columns.str.strip()

            # Crear diccionario de evaluadores
            evaluador_dict = pd.Series(
                df_asignaciones["EVALUADOR"].values,
                index=df_asignaciones["EXPEDIENTE"]
            ).to_dict()

            # Añadir columna EVALASIGN
            df_consolidado["EVALASIGN"] = df_consolidado.apply(
                lambda row: calcular_evalasign(row, evaluador_dict), axis=1
            )

            # Crear un nuevo archivo sin sobrescribir el original
            guardar_como_tabla_nueva(output_path, df_consolidado, f"BASE_{tipo}")

        except Exception as e:
            print(f"Error al procesar {tipo}: {e}")

# Lógica para calcular la columna EVALASIGN
def calcular_evalasign(row, evaluador_dict):
    if row["Evaluado"] == "NO":
        return evaluador_dict.get(row["NumeroTramite"], None)
    elif row["Evaluado"] == "SI" and row["Pre_Concluido"] == "SI":
        return row.get("OperadorPre", None)
    return None

# Función para guardar como un nuevo archivo con formato tabla
def guardar_como_tabla_nueva(archivo, df, tabla_nombre):
    try:
        wb = Workbook()
        ws = wb.active

        # Renombrar la pestaña
        ws.title = tabla_nombre

        # Escribir los datos
        for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), start=1):
            for c_idx, value in enumerate(row, start=1):
                ws.cell(row=r_idx, column=c_idx, value=value)

        # Crear tabla
        table_ref = f"A1:{chr(64 + df.shape[1])}{df.shape[0] + 1}"
        table = Table(displayName=tabla_nombre, ref=table_ref)
        style = TableStyleInfo(
            name="TableStyleMedium9", showFirstColumn=False,
            showLastColumn=False, showRowStripes=True, showColumnStripes=True
        )
        table.tableStyleInfo = style
        ws.add_table(table)

        # Guardar como un nuevo archivo
        wb.save(archivo)
        print(f"Nuevo archivo creado: {archivo}")
    except Exception as e:
        print(f"Error al guardar nuevo archivo: {e}")

# Ejecutar
procesar_cruce()
