import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.table import Table, TableStyleInfo
from file_utils import confirmar_sobrescritura

# Configuración de rutas relativas
current_dir = os.path.dirname(os.path.abspath(__file__))
carpeta_manejo = os.path.join(current_dir, "manejo_reportes")
carpeta_descargas = os.path.join(current_dir, "descargas")

archivos_asignaciones = {
    "CCM": os.path.join(carpeta_manejo, "CCM", "ASIGNACIONES.xlsx"),
    "PRR": os.path.join(carpeta_manejo, "PRR", "ASIGNACIONES.xlsx"),
    "CCM-ESP": os.path.join(carpeta_manejo, "CCM", "ASIGNACIONES.xlsx"),
    "SOL": os.path.join(carpeta_manejo, "SOL", "ASIGNACIONES.xlsx")
}

consolidados = {
    "CCM": os.path.join(carpeta_descargas, "CCM", "Consolidado_CCM.xlsx"),
    "PRR": os.path.join(carpeta_descargas, "PRR", "Consolidado_PRR.xlsx"),
    "CCM-ESP": os.path.join(carpeta_descargas, "CCM-ESP", "Consolidado_CCM-ESP.xlsx"),
    "SOL": os.path.join(carpeta_descargas, "SOL", "Consolidado_SOL.xlsx")
}

consolidados_filtrados = {
    "CCM": os.path.join(carpeta_descargas, "CCM", "Consolidado_Filtrado_CCM.xlsx"),
    "PRR": os.path.join(carpeta_descargas, "PRR", "Consolidado_Filtrado_PRR.xlsx"),
    "CCM-ESP": os.path.join(carpeta_descargas, "CCM-ESP", "Consolidado_Filtrado_CCM.xlsx"),
}

hojas_evaluadores = {
    "CCM": "CONSOLIDADO_CCM_X_EVAL",
    "PRR": "CONSOLIDADO_PRR_X_EVAL",
    "CCM-ESP": "CONSOLIDADO_CCM_X_EVAL",
    "SOL": "CONSOLIDADO_SOL_X_EVAL"
}

def procesar_cruces_combinados():
    for tipo in consolidados.keys():
        try:
            consolidado_path = consolidados[tipo]
            asignaciones_path = archivos_asignaciones[tipo]
            consolidado_filtrado_path = consolidados_filtrados.get(tipo)
            hoja_evaluador = hojas_evaluadores[tipo]
            output_path = consolidado_path.replace(".xlsx", "_CRUZADO.xlsx")

            df_consolidado = pd.read_excel(consolidado_path)
            df_asignaciones = pd.read_excel(asignaciones_path, sheet_name=hoja_evaluador)
            df_asignaciones.columns = df_asignaciones.columns.str.strip()

            df_asignaciones["EXPEDIENTE"] = df_asignaciones["EXPEDIENTE"].astype(str)
            evaluador_dict = pd.Series(df_asignaciones["EVALUADOR"].values, index=df_asignaciones["EXPEDIENTE"]).to_dict()

            df_consolidado["EVALASIGN"] = df_consolidado.apply(lambda row: calcular_evalasign(row, evaluador_dict), axis=1)

            if tipo != "SOL" and consolidado_filtrado_path and os.path.exists(consolidado_filtrado_path):
                df_filtrado = pd.read_excel(consolidado_filtrado_path)
                df_filtrado.columns = df_filtrado.columns.str.strip()
                df_filtrado["EXPEDIENTE"] = df_filtrado["EXPEDIENTE"].astype(str)
                
                for col in ["ESTADO", "DESCRIPCION", "FECHA DE TRABAJO"]:
                    if df_filtrado[col].dtype == 'object':
                        df_filtrado[col] = pd.Categorical(df_filtrado[col], ordered=False)

                if "FECHA DE TRABAJO" in df_filtrado.columns:
                    df_filtrado["FECHA DE TRABAJO"] = pd.to_datetime(df_filtrado["FECHA DE TRABAJO"], errors='coerce')
                
                df_filtrado = df_filtrado.sort_values(["EXPEDIENTE", "FECHA DE TRABAJO"], ascending=[True, False])
                df_filtrado = df_filtrado.drop_duplicates(subset="EXPEDIENTE", keep="first")

                df_consolidado = pd.merge(
                    df_consolidado,
                    df_filtrado[["EXPEDIENTE", "ESTADO", "DESCRIPCION", "FECHA DE TRABAJO"]],
                    left_on="NumeroTramite",
                    right_on="EXPEDIENTE",
                    how="left"
                ).drop(columns=["EXPEDIENTE"], errors="ignore")

            guardar_como_tabla_nueva(output_path, df_consolidado, f"BASE_{tipo}")

        except Exception as e:
            print(f"Error al procesar {tipo}: {str(e)}")

def calcular_evalasign(row, evaluador_dict):
    numero_tramite = str(row.get("NumeroTramite", ""))
    if row["Evaluado"] == "NO":
        return evaluador_dict.get(numero_tramite, None)
    elif row["Evaluado"] == "SI" and row["Pre_Concluido"] == "SI":
        return row.get("OperadorPre", None)
    return None

def guardar_como_tabla_nueva(archivo, df, tabla_nombre):
    wb = Workbook()
    ws = wb.active
    ws.title = tabla_nombre

    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), start=1):
        for c_idx, value in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=value)

    table_ref = f"A1:{chr(64 + df.shape[1])}{df.shape[0] + 1}"
    table = Table(displayName=tabla_nombre, ref=table_ref)
    style = TableStyleInfo(
        name="TableStyleMedium9", showFirstColumn=False,
        showLastColumn=False, showRowStripes=True, showColumnStripes=True
    )
    table.tableStyleInfo = style
    ws.add_table(table)
    wb.save(archivo)

if __name__ == "__main__":
    procesar_cruces_combinados()